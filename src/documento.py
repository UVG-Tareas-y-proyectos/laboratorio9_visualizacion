"""Genera el informe académico final en PDF en la raíz del repositorio.

Uso: python src/documento.py
La exportación a PDF utiliza Word en Windows o LibreOffice cuando esté disponible.
"""
from pathlib import Path
import json
import os
import shutil
import subprocess

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Mm
from docx.opc.constants import RELATIONSHIP_TYPE
from pypdf import PdfReader, PdfWriter

ROOT=Path(__file__).resolve().parents[1]
URL='https://github.com/UVG-Tareas-y-proyectos/laboratorio9_visualizacion'


def exportar_pdf(docx_path):
    if os.name=='nt':
        subprocess.run(['powershell.exe','-NoProfile','-ExecutionPolicy','Bypass','-File',
                        str(ROOT/'src/exportar_pdf.ps1'),'-InputPath',str(docx_path)],check=True)
    else:
        soffice=shutil.which('soffice') or shutil.which('libreoffice')
        if not soffice:
            raise RuntimeError('El DOCX se guardó. Para exportar PDF se necesita Word en Windows o LibreOffice.')
        subprocess.run([soffice,'--headless','--convert-to','pdf','--outdir',str(docx_path.parent),str(docx_path)],check=True)


def main():
    doc=Document()
    section=doc.sections[0]
    section.page_width=Mm(210);section.page_height=Mm(297)
    section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Inches(1)
    cover=[('UNIVERSIDAD DEL VALLE DE GUATEMALA',12,True,5,1.0),
           ('DATA SCIENCE - 10',12,False,5,1.0),
           ('Ing. Erick Marroquin',12,False,34,1.0),
           ('',12,False,60,1.0),
           ('Laboratorio 9. Visualización estática',15,True,85,1.5),
           ('Ihan Marroquin - 23108\nDiego Patzan - 23525\n\n',12,False,82,1.5),
           ('GUATEMALA, 18 de octubre del 2026',12,True,0,1.0)]
    for i,(text,size,bold,after,line) in enumerate(cover):
        par=doc.add_paragraph();par.alignment=WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after=Pt(after);par.paragraph_format.line_spacing=line
        run=par.add_run(text);run.font.name='Times New Roman';run.font.size=Pt(size);run.bold=bold
        if i==1:run.italic=True
        if i==3:run.add_picture(str(ROOT/'src/recursos/logo_uvg.png'),width=Inches(2.09375))
    named_styles={s.name:s for s in doc.styles}
    for name,size,bold in [('Normal',12,False),('Heading 1',14,True),('Heading 2',12,True),('Caption',9.5,False)]:
        style=named_styles.get(name)
        if style is None and name=='Normal':style=named_styles.get('normal')
        if style is None:style=doc.styles.add_style(name,1)
        style.font.name='Times New Roman';style.font.size=Pt(size);style.font.bold=bold
        style.font.color.rgb=RGBColor(0,0,0)
    normal=named_styles.get('Normal')
    if normal is None:normal=named_styles['normal']
    normal.paragraph_format.line_spacing=1.07
    normal.paragraph_format.space_after=Pt(6)
    normal.paragraph_format.widow_control=True
    doc.add_page_break()
    def p(s,small=False):
        par=doc.add_paragraph(style=normal)
        par.alignment=(WD_ALIGN_PARAGRAPH.CENTER if small else
                       WD_ALIGN_PARAGRAPH.LEFT if '\n' in s else WD_ALIGN_PARAGRAPH.JUSTIFY)
        run=par.add_run(s);run.font.name='Times New Roman';run.font.size=Pt(9.5 if small else 12)
        par.paragraph_format.line_spacing=1.0 if small else 1.07
        par.paragraph_format.space_after=Pt(6)
        return par
    def h(s,level=1):
        par=doc.add_paragraph(s,style=f'Heading {level}')
        for run in par.runs:
            run.font.name='Times New Roman';run.font.size=Pt(14 if level==1 else 12)
            run.font.bold=True;run.font.color.rgb=RGBColor(0,0,0)
        par.alignment=WD_ALIGN_PARAGRAPH.LEFT
        par.paragraph_format.space_before=Pt(12 if level==1 else 8)
        par.paragraph_format.space_after=Pt(6)
        par.paragraph_format.keep_with_next=True
        par.paragraph_format.keep_together=True
        par.paragraph_format.line_spacing=1.0
    def table(headers,rows,widths):
        t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
        usable=doc.sections[0].page_width-doc.sections[0].left_margin-doc.sections[0].right_margin
        scaled=[int(usable*v/sum(widths)) for v in widths]
        for col,width in zip(t.columns,scaled):col.width=width
        border=OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'B7B7B7');border.append(el)
        t._tbl.tblPr.append(border)
        repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for i,values in enumerate([headers,*rows]):
            row=t.rows[0] if i==0 else t.add_row()
            cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
            for j,(cell,value) in enumerate(zip(row.cells,values)):
                cell.width=scaled[j];cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cp=cell._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
                for edge,val in [('top','60'),('bottom','60'),('left','90'),('right','90')]:
                    el=OxmlElement('w:'+edge);el.set(qn('w:w'),val);el.set(qn('w:type'),'dxa');margins.append(el)
                cp.append(margins)
                if i==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7E6E6');cp.append(shade)
                par=cell.paragraphs[0];par.alignment=WD_ALIGN_PARAGRAPH.LEFT
                par.paragraph_format.space_before=Pt(0);par.paragraph_format.space_after=Pt(0)
                par.paragraph_format.line_spacing=1.0
                run=par.add_run(str(value));run.font.name='Times New Roman';run.font.size=Pt(10.5);run.bold=i==0
        doc.add_paragraph().paragraph_format.space_after=Pt(0)
    def image(name,caption,width=505):
        par=doc.add_paragraph();par.alignment=WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after=Pt(2);par.paragraph_format.keep_with_next=True
        par.add_run().add_picture(str(ROOT/'results/figures'/f'{name}.png'),width=Inches(6.1))
        p(caption,True)
    def page():
        # Flujo continuo para evitar páginas casi vacías al usar letra de 12 puntos.
        pass
    def hyperlink(label,target):
        par=doc.add_paragraph();par.paragraph_format.space_after=Pt(0);par.paragraph_format.line_spacing=1.0
        rel=doc.part.relate_to(target,RELATIONSHIP_TYPE.HYPERLINK,is_external=True)
        link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rel)
        run=OxmlElement('w:r');pr=OxmlElement('w:rPr')
        font=OxmlElement('w:rFonts');font.set(qn('w:ascii'),'Times New Roman');font.set(qn('w:hAnsi'),'Times New Roman');pr.append(font)
        size=OxmlElement('w:sz');size.set(qn('w:val'),'21');pr.append(size)
        color=OxmlElement('w:color');color.set(qn('w:val'),'0000FF');pr.append(color)
        run.append(pr);text=OxmlElement('w:t');text.text=label;run.append(text);link.append(run);par._p.append(link)
    tb=ROOT/'results/tables';a=json.loads((tb/'auditoria.json').read_text(encoding='utf-8'))
    k=a['kpis'];v=a['validacion']
    res=pd.read_csv(tb/'resumen_series_lab1.csv');models=pd.read_csv(tb/'comparacion_modelos_lab1.csv')
    win=pd.read_csv(tb/'ganadores_lab1.csv')
    from visualizacion import SHORT,short_model

    hyperlink('Repositorio:',URL)
    hyperlink('Infografia:',URL+'/blob/main/entrega/Infografia_turismo_Guatemala.png')
    h('Explicación de las decisiones de diseño')
    p('La infografía comunica cómo cambió el ingreso mensual de viajeros entre enero de 2009 y junio de 2026 y qué límites tienen los modelos entrenados antes de la recuperación. La conclusión central es que el flujo se recuperó del desplome de 2020, pero el pronóstico del Lab 1 quedó persistentemente por debajo de lo observado. Por eso proponemos decisiones operativas apoyadas en la estacionalidad y una actualización del modelo antes de usarlo para presupuestos.')
    h('1 Audiencia y recorrido de lectura')
    p('Nos dirigimos a una persona del INGUAT que planifica campañas, personal y capacidad. Tiene pocos minutos para leer y no necesita conocer la formulación estadística. Usamos un titular informativo, cifras grandes, gráficos con unidades y recomendaciones vinculadas a evidencia. Los términos técnicos que no caben en el lienzo se definen en este documento.')
    table(['Bloque del enunciado','Representación final y propósito'],[
        ['Titular y cifras','Titular sobre recuperación y pronóstico; cuatro indicadores dan escala al problema.'],
        ['Panorama histórico','Línea mensual con eventos de contexto y banda del período usado para medir la caída.'],
        ['Estacionalidad vías y mercados','Barras por mes, rankings horizontales y matriz de seis series.'],
        ['Modelos y pronóstico','Tabla de ocho algoritmos sobre el mismo conjunto de prueba y predicción con intervalo del 95%.'],
        ['Hallazgos y recomendaciones','Tres acciones concretas para campañas, segmentación y actualización predictiva.']],[145,360])
    p('El recorrido de arriba hacia abajo va del contexto a la decisión. El titular y los indicadores tienen mayor tamaño; la historia precede al calendario; la comparación estadística precede a las recomendaciones. Los cinco bloques del enunciado están presentes: el titular no lleva número y los otros cuatro se identifican del 01 al 04.')
    h('2 Alcance de los datos')
    p('Usamos la Base de Migración 2009-2026 INGUAT/IGM facilitada en el curso. Conservamos Turista y Excursionista, como en el Lab 1; excluimos Crucerista y Viajero para mantener el mismo alcance. Son registros agregados de ingresos de viajeros, no personas únicas. El período tiene 210 observaciones mensuales; 2026 solo incluye enero a junio. Las cifras no se presentan como una nueva publicación oficial de las instituciones.')
    p('La variable País cambia de nivel de detalle después de 2023, según la advertencia del análisis previo. Guatemala se conserva como categoría registrada en la base y no se interpreta automáticamente como turismo interno, nacionalidad ni mercado internacional independiente. Esta limitación impide convertir el ranking acumulado en una recomendación de inversión por país sin validación adicional.')
    page()

    h('3 Historia y estacionalidad')
    p('La línea es apropiada para observaciones mensuales ordenadas: conserva la secuencia, muestra crecimiento, caída y recuperación sin introducir agregaciones anuales engañosas. El eje inicia en cero. Las anotaciones de H1N1 y Volcán de Fuego ubican hechos que plantea el análisis previo; no constituyen pruebas de impacto o ausencia de impacto. Sustituimos la afirmación absoluta de las imágenes recibidas por una lectura descriptiva.')
    image('01_historia_eventos','Figura 1. La banda abril 2020-marzo 2021 es una ventana de medición, no la duración completa de la pandemia.',490)
    p('Para el calendario calculamos 100 × (viajeros del mes / promedio mensual de su año - 1), y luego promediamos por mes del calendario dando el mismo peso a cada año. Excluimos 2020 y 2021 por la ruptura y 2026 porque está incompleto. Quedan 15 años completos: 2009-2019 y 2022-2025. Las barras parten de cero y separan valores positivos, negativos y diciembre mediante color y etiquetas numéricas.')
    image('02_estacionalidad','Figura 2. Diciembre está 35.34% sobre el promedio anual; febrero -17.83% y septiembre -16.89%. Valores del lienzo redondeados.',490)
    p('El patrón coincide con una mayor movilidad al cierre de año; la base no permite demostrar que aguinaldo o fiestas sean su causa. Tampoco atribuimos los pequeños aumentos de marzo, abril o julio a Semana Santa o Bono 14 sin identificar fechas y contrastes adicionales. La estacionalidad descriptiva de años completos y la fuerza STL de entrenamiento responden preguntas distintas y no deben confundirse.')
    page()

    h('4 Cifras vías mercados y comparación de series')
    table(['Indicador','Cálculo y resultado'],[
        ['Promedio mensual',f'Suma de los 210 meses / 210 = {k["promedio_mensual"]:,.2f}; se muestra 222,438.'],
        ['Máximo mensual',f'{k["maximo_mensual"]:,.2f} en diciembre de 2019; se muestra 449,114.'],
        ['Caída en pandemia','1 - promedio abr 2020-mar 2021 / promedio ene-dic 2019 = 89.6958%; se muestra 89.7%.'],
        ['La Aurora acumulado',f'{k["la_aurora_acumulado"]:,.2f} en el período; se muestra 18.99 millones.']],[140,365])
    p('Los rankings usan barras horizontales con origen cero porque comparan categorías sin orden natural y nombres largos. La vía terrestre acumula 27.59 millones, la aérea 19.02 millones y la marítima 100,241 ingresos comparables. El volumen marítimo excluye Crucerista y por tanto no representa todo el turismo de cruceros. El Salvador, Guatemala y Estados Unidos son los tres mayores valores acumulados de País: 14.08, 13.87 y 6.97 millones, con las restricciones ya descritas.')
    h('Definiciones de las cuatro medidas',2)
    p('Estacionalidad: max(0, 1 - Var(residuo) / Var(estacional + residuo)) de una STL robusta sobre log(1+y), período 12, usando solo entrenamiento. Crecimiento: pendiente de una recta sobre log(1+y) entre enero de 2022 y junio de 2026; es una medida mensual en escala logarítmica, no un porcentaje directo. Volatilidad: desviación estándar poblacional de las primeras diferencias de log(1+y) de toda la serie. Caída: reducción del promedio abril 2020-marzo 2021 frente al promedio de 2019.')
    rows=[]
    for _,r in res[res.Serie.ne('Total mensual')].iterrows():
        rows.append([SHORT[r.Serie],f'{r["Fuerza estacional"]:.3f}',f'{r["Crecimiento recuperación"]:+.4f}',f'{r["Volatilidad"]:.3f}',f'{r["Caída pandemia"]*100:.1f}%'])
    table(['Serie','Fuerza STL','Pendiente','Volatilidad','Caída'],rows,[171,80,85,89,80])
    p('La tabla visual permite responder las cuatro preguntas sin combinar magnitudes incompatibles en un único índice. El ocre resalta el mayor valor dentro de regiones y dentro de fronteras. Entre regiones, Europa presenta mayor fuerza estacional y caída; América del Norte, mayor crecimiento y volatilidad. Entre fronteras, San Cristóbal presenta mayor estacionalidad y crecimiento; Valle Nuevo, mayor volatilidad y caída. En las seis series, los máximos globales son San Cristóbal para estacionalidad y crecimiento, América del Norte para volatilidad y Valle Nuevo para caída.')
    h('Recomendaciones derivadas de la evidencia',2)
    p('Primero, preparar capacidad para diciembre y probar campañas en febrero y septiembre con metas medibles de demanda. Segundo, segmentar las decisiones por región y frontera: las fuertes caídas históricas de Europa y Valle Nuevo justifican revisar su recuperación con datos actuales, antes de reasignar presupuesto. Tercero, reentrenar y validar por ventanas móviles con datos recientes; el corte de entrenamiento en pandemia limita el pronóstico. Los resultados históricos orientan estas revisiones y no prueban por sí solos dónde invertir hoy.')
    page()

    h('5 Comparación de modelos')
    p('Conservamos la división temporal del Lab 1: enero 2009-marzo 2021 para entrenamiento (147 meses) y abril 2021-junio 2026 para prueba (63 meses). MAE es el promedio del error absoluto; RMSE penaliza con más fuerza los errores grandes. Ambos se expresan en viajeros y deben compararse entre algoritmos de la misma serie y horizonte. No se usan para concluir que una región es más predecible que otra con diferente volumen.')
    total=models[models.Serie.eq('Total mensual')].sort_values(['RMSE','MAE'])
    table(['Modelo total mensual','MAE viajeros','RMSE viajeros'],[
        [short_model(r.Modelo),f'{r.MAE:,.2f}',f'{r.RMSE:,.2f}'] for _,r in total.iterrows()],[279,113,113])
    p('El ganador guardado del total es ARIMA(2,1,0), por menor RMSE, con MAE como segundo criterio de orden. Los valores de esta tabla se leen del CSV original del Lab 1. El modelo rotulado como respaldo aditivo usa una regresión Ridge con tendencia, cambios de pendiente y términos de Fourier; no es Prophet oficial. SARIMA y las variantes ARIMA se incluyen tal como fueron evaluadas. AIC y BIC no se usan para ordenar familias diferentes ni series diferentes; están disponibles en las tablas fuente.')
    h('Ganadores originales de las siete series',2)
    table(['Serie','Modelo ganador','MAE','RMSE'],[
        [SHORT[r.Serie],short_model(r.Modelo),f'{r.MAE:,.0f}',f'{r.RMSE:,.0f}'] for _,r in win.iterrows()],[160,173,86,86])
    p('Ningún método gana en todas las series. El Lab 1 guardado elige ARIMA(2,1,0) para total, América del Centro y América del Norte; ARIMA(0,1,2) para La Aurora; exponencial simple para Europa y San Cristóbal; Holt-Winters aditivo para Valle Nuevo. Estas filas corrigen el avance, que utilizó candidatos y cifras de otra comparación. La versión entregada no mezcla esos resultados.')
    h('Alcance de esta elección',2)
    p('El conjunto de prueba se usa para comparar modelos y elegir el ganador histórico. Eso no equivale a una evaluación final independiente de un modelo ya elegido. Además, casi todo el horizonte de prueba pertenece a la recuperación y el entrenamiento termina en un nivel muy bajo. Antes de recomendar un pronóstico operativo corresponde validar por ventanas móviles y reservar una evaluación futura, en lugar de extrapolar esta tabla hasta diciembre de 2026.')
    page()

    h('6 Predicción e incertidumbre')
    p(f'Reajustamos únicamente el ganador del total con la especificación original: SARIMAX sobre log(1+y), orden (2,1,0), sin tendencia y sin imponer estacionariedad ni invertibilidad. El ajuste convergió. La validación reproduce MAE {v["MAE_recalculado"]:,.2f} y RMSE {v["RMSE_recalculado"]:,.2f}, con diferencias inferiores a 0.00001 viajeros frente al CSV conservado.')
    image('03_prediccion_ic95','Figura 3. Serie observada, predicción del ganador del Lab 1 e intervalo del 95% rotulado Estimación. La banda excede el eje visible.',490)
    p(f'Los límites se calculan en escala logarítmica y se transforman con exp(x)-1, con piso cero. La curva puntual mantiene la convención del Lab 1: exp(media logarítmica)-1; corresponde a una mediana condicional, no a la media aritmética con ajuste de sesgo. El primer intervalo es aproximadamente [{v["primer_inferior95"]:,.0f}, {v["primer_superior95"]:,.0f}]; el límite superior final alcanza {v["ultimo_superior95"]/1e6:.2f} millones.')
    image('04_prediccion_intervalo_completo','Figura 4. Vista complementaria en escala logarítmica que muestra la amplitud completa del intervalo; no se usa para comparar alturas como proporciones lineales.',490)
    p('El intervalo es predictivo y condicionado a los supuestos del modelo; en el lienzo se conserva el rótulo 95% que pide el enunciado. Los 63 valores observados caen dentro de él, pero una cobertura de 100% con amplitud extrema no demuestra una buena predicción. El intervalo no incorpora explícitamente incertidumbre sobre cambios estructurales. El eje de la infografía se limita a 650 mil para leer la serie; una nota declara el recorte y el límite completo, que se muestra aquí. Todo el gráfico es retrospectivo sobre prueba, no una proyección posterior a junio de 2026.')
    page()

    h('7 Decisiones de diseño y teoría del color')
    p('Conservamos la dirección visual propuesta en el avance: una familia sans serif, fondo crema, línea histórica verde y un acento cálido para dirigir la atención. Verde y turquesa forman una relación cercana en el círculo cromático; el ocre introduce contraste de temperatura. Las asociaciones con naturaleza, agua y maíz son decisiones de diseño contextual y no propiedades estadísticas de los datos.')
    table(['Color','Función semántica'],[
        ['Verde selva #124734','Serie observada, barras positivas y jerarquía visual.'],
        ['Turquesa #2E9CA6','Desviaciones mensuales negativas; se acompañan de signo y etiqueta.'],
        ['Ocre #C27C00','Pico de diciembre, banda histórica, ganador y predicción. El texto pequeño usa un tono más oscuro de la misma familia.'],
        ['Crema #FAF6EC y gris #2B2B2B','Fondo neutro y texto principal, respectivamente.']],[180,325])
    p('El color no es el único canal: observación y predicción se diferencian con línea continua y discontinua; los valores negativos llevan signo; los ganadores están en negrita; las bandas están rotuladas. Evitamos afirmar accesibilidad completa para todos los tipos de daltonismo sin pruebas. El acento de texto se oscurece para mejorar el contraste sobre crema. Se evita depender del turquesa para leer cifras pequeñas.')
    p('Las barras parten de cero, se rotulan unidades y se eliminan 3D, sombras, fondos saturados y leyendas lejanas. El único recorte deliberado es el eje del intervalo predictivo y se declara junto al gráfico. Las reglas de separación son discretas; los gráficos conservan espacio para ejes y etiquetas. Las tarjetas de indicadores se construyen con tipografía y espacio, sin cajas decorativas.')
    h('8 Correcciones frente a lo recibido')
    table(['Elemento','Decisión verificable'],[
        ['Ganador del total','Se conserva ARIMA(2,1,0) del CSV del Lab 1. ARIMA(0,1,1) del script recibido es otro candidato y no reemplaza el análisis original.'],
        ['Intervalo final','10.63 millones con el ganador original; 20.64 millones pertenecen al candidato del avance.'],
        ['Estacionalidad','Se excluye 2026 incompleto además de 2020-2021. Enero queda -4%, febrero -18% y marzo +5% al redondear.'],
        ['Eventos y fiestas','Se reemplazan afirmaciones causales por observaciones descriptivas y contexto.'],
        ['Respaldo aditivo','Se identifica expresamente como respaldo del Lab 1 y no como Prophet oficial.']],[132,373])
    h('9 Formatos del documento y la infografía')
    p('El PNG tiene 3507 × 4960 píxeles, superior al mínimo 1200 × 2400. El enunciado mezcla dos tamaños de impresión distintos: A3 mide 297 × 420 mm y 11 × 17 pulgadas mide 279.4 × 431.8 mm. Se entregan dos PDF separados de una página, uno por cada tamaño, manteniendo el mismo contenido. El informe se entrega únicamente en PDF en la raíz del repositorio. La infografía completa ocupa una página independiente al final de este documento.')
    page()

    h('10 Reproducción y archivos de trabajo')
    p('El procesamiento es local. Los scripts de src generan tablas, gráficas, infografía e informe y verifican dimensiones, páginas y métricas. results/tables/auditoria.json registra versiones y SHA-256 de los insumos. Solo se redondea al presentar resultados.')
    p('Para reproducir los resultados en Windows con Python 3.12, ejecutar los siguientes comandos desde la raíz del repositorio. La generación del informe PDF utiliza Microsoft Word en Windows o LibreOffice en otros sistemas.')
    p('python -m venv .venv\n.venv\\Scripts\\Activate.ps1\npip install -r requirements.txt\npython src/analisis.py\npython src/visualizacion.py\npython src/documento.py\npython src/verificar_entrega.py')
    doc.core_properties.title='Laboratorio 9 Visualización estática'
    doc.core_properties.author='Ihan Marroquin; Diego Patzan'
    doc.core_properties.subject='Explicación de las decisiones de diseño y reproducción del laboratorio'
    work=ROOT/'tmp/documento';work.mkdir(parents=True,exist_ok=True)
    intermediate=work/'Informe_Laboratorio9.docx'
    doc.save(intermediate)
    exportar_pdf(intermediate)
    # La lámina A3 usa una página completa y conserva la nitidez de la infografía.
    writer=PdfWriter()
    writer.append(PdfReader(intermediate.with_suffix('.pdf')))
    writer.append(PdfReader(ROOT/'entrega/Infografia_turismo_Guatemala_A3.pdf'))
    writer.add_metadata({'/Title':'Laboratorio 9 Visualización estática',
                         '/Author':'Ihan Marroquin; Diego Patzan'})
    out=ROOT/'Informe_Laboratorio9.pdf'
    with out.open('wb') as stream:writer.write(stream)
    print(out)


if __name__=='__main__':
    main()
